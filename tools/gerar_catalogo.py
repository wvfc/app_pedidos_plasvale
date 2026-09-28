#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gera app/src/main/assets/catalogo.json a partir da planilha de pedidos da Plasvale.

Uso:
    pip install openpyxl
    python3 tools/gerar_catalogo.py "FOLHA DE PEDIDOS TABELA 130 ... .xlsx"

A planilha precisa conter as abas:
  - "Tabela 130 "   -> cabeçalho na linha 5, produtos a partir da linha 6
                       (A=Referência, B=Descrição, C=EMB., D..N=preços por prazo)
  - "FOLHA DE PEDIDO " -> coluna U com a lista de UFs usada na validação de dados
  - "BASE CONTÁBIL" -> alíquotas de ST por NCM/UF (colunas A..K)
"""
import json
import sys
import unicodedata

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("Instale a dependência: pip install openpyxl")

ABA_TABELA = "Tabela 130 "
ABA_PEDIDO = "FOLHA DE PEDIDO "
ABA_BASE = "BASE CONTÁBIL"

LINHA_CABECALHO = 5
PRIMEIRA_LINHA_PRODUTO = 6
COLUNAS_ST = ["MG", "PR", "RJ", "RS", "SP"]

AVISOS = [
    "Região de SP desconsiderar a ST.",
    "Tabela não válida para clientes do Simples Nacional.",
    "Desconto de 20% pode ser aplicado a todos os itens.",
    "Pedido mínimo R$ 3.000,00: pequeno e médio varejo.",
    "Pedido mínimo R$ 5.000,00: supermercados e grandes redes.",
]


def normaliza_aba(wb, desejada):
    """Encontra a aba ignorando espaços e acentos."""
    def chave(t):
        t = unicodedata.normalize("NFD", t)
        t = "".join(c for c in t if unicodedata.category(c) != "Mn")
        return t.strip().upper()
    alvo = chave(desejada)
    for nome in wb.sheetnames:
        if chave(nome) == alvo:
            return wb[nome]
    raise SystemExit(f"Aba não encontrada: {desejada} (abas: {wb.sheetnames})")


def numero(valor):
    if isinstance(valor, str) and valor.strip().endswith("%"):
        return round(float(valor.replace("%", "").replace(",", ".")) / 100, 4)
    return float(valor) if isinstance(valor, (int, float)) else 0.0


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    caminho = sys.argv[1]
    destino = sys.argv[2] if len(sys.argv) > 2 else "app/src/main/assets/catalogo.json"

    wb = openpyxl.load_workbook(caminho, data_only=True)
    tabela = normaliza_aba(wb, ABA_TABELA)

    prazos = []
    coluna = 4
    while True:
        titulo = tabela.cell(row=LINHA_CABECALHO, column=coluna).value
        if titulo is None or str(titulo).strip() in ("MG", "PR", "RJ", "RS", "SP", "IPI"):
            break
        prazos.append(str(titulo).strip())
        coluna += 1

    produtos = []
    linha = PRIMEIRA_LINHA_PRODUTO
    vazias = 0
    while vazias < 30:
        referencia = tabela.cell(row=linha, column=1).value
        if referencia in (None, ""):
            vazias += 1
            linha += 1
            continue
        vazias = 0
        precos = {}
        for i, prazo in enumerate(prazos):
            valor = tabela.cell(row=linha, column=4 + i).value
            precos[prazo] = round(float(valor), 2) if isinstance(valor, (int, float)) else None
        embalagem = tabela.cell(row=linha, column=3).value
        produtos.append({
            "referencia": str(referencia).strip(),
            "descricao": str(tabela.cell(row=linha, column=2).value or "").strip(),
            "embalagem": int(embalagem) if isinstance(embalagem, (int, float)) else 0,
            "precos": precos,
        })
        linha += 1

    pedido = normaliza_aba(wb, ABA_PEDIDO)
    ufs = []
    for r in range(2, 60):
        uf = pedido.cell(row=r, column=21).value  # coluna U
        if uf:
            ufs.append(str(uf).strip())

    base = normaliza_aba(wb, ABA_BASE)
    st = []
    for r in range(3, 60):
        ncm = base.cell(row=r, column=1).value
        if ncm in (None, ""):
            continue
        st.append({
            "ncm": str(ncm).strip(),
            "revenda": {uf: numero(base.cell(row=r, column=2 + i).value) for i, uf in enumerate(COLUNAS_ST)},
            "nacional": {uf: numero(base.cell(row=r, column=7 + i).value) for i, uf in enumerate(COLUNAS_ST)},
        })

    saida = {
        "tabela": "130",
        "referencia": "SETEMBRO 2026",
        "ipiPadrao": 6.5,
        "avisos": AVISOS,
        "prazos": prazos,
        "ufs": ufs,
        "produtos": produtos,
        "stPorNcm": st,
    }
    with open(destino, "w", encoding="utf-8") as arq:
        json.dump(saida, arq, ensure_ascii=False, indent=1)
    print(f"{len(produtos)} produtos, {len(prazos)} prazos, {len(ufs)} UFs -> {destino}")


if __name__ == "__main__":
    main()
